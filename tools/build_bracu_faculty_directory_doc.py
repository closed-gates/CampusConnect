from pathlib import Path
from datetime import date
from urllib.parse import urljoin

from lxml import html
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "artifacts_faculty_source.html"
OUTPUT_DIR = ROOT / "artifacts"
OUTPUT = OUTPUT_DIR / "BRACU_CSE_Faculty_and_Staff_Directory.docx"
BASE_URL = "https://cse.bracu.ac.bd"


def clean(text):
    return " ".join((text or "").split())


def extract_records():
    tree = html.fromstring(SOURCE.read_bytes())
    records = []
    current_section = "Department Leadership"

    nodes = tree.xpath(
        "//*[self::h2[contains(concat(' ', normalize-space(@class), ' '), ' facList-title ')] "
        "or self::div[contains(concat(' ', normalize-space(@class), ' '), ' fac-card ')]]"
    )
    for node in nodes:
        if node.tag == "h2":
            current_section = clean(node.text_content())
            continue

        name_nodes = node.xpath(
            ".//a[contains(@href, '/faculty_profile/')]/p[contains(@class, 'text-black')]"
        )
        if not name_nodes:
            continue
        name = clean(name_nodes[0].text_content())
        profile = name_nodes[0].getparent().get("href", "")
        profile = urljoin(BASE_URL, profile)

        paragraphs = [clean(p.text_content()) for p in node.xpath(".//p")]
        paragraphs = [p for p in paragraphs if p]
        email = next((p for p in reversed(paragraphs) if "@" in p), "")
        titles = [p for p in paragraphs if p not in {name, email}]
        position = " | ".join(dict.fromkeys(titles))

        records.append(
            {
                "section": current_section,
                "name": name,
                "position": position,
                "email": email,
                "profile": profile,
            }
        )
    return records


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="D9D9D9", size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), size)
        el.set(qn("w:color"), color)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def keep_row_together(row):
    for cell in row.cells:
        for paragraph in cell.paragraphs:
            p_pr = paragraph._p.get_or_add_pPr()
            keep = OxmlElement("w:keepLines")
            p_pr.append(keep)


def add_hyperlink(paragraph, text, url, color="0563C1"):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    run_props = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    run_props.append(c)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    run_props.append(underline)
    run.append(run_props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def style_run(run, size=10.5, bold=False, color="222222"):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def build_document(records):
    OUTPUT_DIR.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Aptos")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(34, 34, 34)
    normal.paragraph_format.space_after = Pt(5)

    for style_name, size in (("Title", 25), ("Heading 1", 16), ("Heading 2", 12)):
        style = styles[style_name]
        style.font.name = "Aptos Display" if style_name != "Heading 2" else "Aptos"
        style._element.rPr.rFonts.set(qn("w:ascii"), style.font.name)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), style.font.name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.add_run("BRACU CSE Faculty and Staff Directory")

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Department of Computer Science and Engineering  |  BRAC University")
    style_run(run, 11, True, "166B73")

    intro = doc.add_paragraph()
    intro.paragraph_format.space_after = Pt(10)
    intro.add_run(
        "This directory organizes the names, positions, email addresses, and profile links "
        "published on the BRACU CSE Faculty and Staff page. Records are grouped according "
        "to the categories shown on the source page."
    )

    counts = {}
    for record in records:
        counts[record["section"]] = counts.get(record["section"], 0) + 1

    summary = doc.add_paragraph()
    summary.paragraph_format.space_after = Pt(14)
    r = summary.add_run(f"Directory snapshot: {len(records)} people across {len(counts)} categories")
    style_run(r, 10.5, True, "166B73")
    retrieved = date.today().strftime("%d %B %Y").lstrip("0")
    r = summary.add_run(f"  |  Retrieved {retrieved}")
    style_run(r, 10, False, "5F6B76")

    source_p = doc.add_paragraph()
    source_p.paragraph_format.space_after = Pt(15)
    style_run(source_p.add_run("Source: "), 9.5, True, "444444")
    add_hyperlink(source_p, "BRACU CSE Faculty and Staff", f"{BASE_URL}/faculty_list")

    by_section = {}
    for record in records:
        by_section.setdefault(record["section"], []).append(record)

    for section_name, people in by_section.items():
        heading = doc.add_paragraph(style="Heading 1")
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(12)
        heading.paragraph_format.space_after = Pt(6)
        heading.add_run(f"{section_name}  {len(people)}")

        table = doc.add_table(rows=1, cols=4)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(0.45)
        table.columns[1].width = Inches(2.55)
        table.columns[2].width = Inches(1.85)
        table.columns[3].width = Inches(2.35)
        set_table_borders(table)

        header = table.rows[0]
        repeat_table_header(header)
        headers = ("No", "Name", "Position", "Email")
        for index, (cell, label) in enumerate(zip(header.cells, headers)):
            cell.width = table.columns[index].width
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "176B73")
            set_cell_margins(cell, 105, 110, 105, 110)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index == 0 else WD_ALIGN_PARAGRAPH.LEFT
            style_run(p.add_run(label), 9.5, True, "FFFFFF")

        for number, person in enumerate(people, 1):
            row = table.add_row()
            keep_row_together(row)
            shade = "F2F8F8" if number % 2 == 0 else "FFFFFF"
            for index, cell in enumerate(row.cells):
                cell.width = table.columns[index].width
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_shading(cell, shade)
                set_cell_margins(cell)

            p = row.cells[0].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            style_run(p.add_run(str(number)), 9.5, False, "555555")

            p = row.cells[1].paragraphs[0]
            add_hyperlink(p, person["name"], person["profile"], "176B73")

            p = row.cells[2].paragraphs[0]
            style_run(p.add_run(person["position"]), 9.5, False, "333333")

            p = row.cells[3].paragraphs[0]
            add_hyperlink(p, person["email"], f"mailto:{person['email']}", "0563C1")

        after = doc.add_paragraph()
        after.paragraph_format.space_after = Pt(2)

    footer = section.footer
    footer_p = footer.paragraphs[0]
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    style_run(
        footer_p.add_run("BRACU CSE Faculty and Staff Directory  |  Public directory snapshot"),
        8.5,
        False,
        "6B7280",
    )

    props = doc.core_properties
    props.title = "BRACU CSE Faculty and Staff Directory"
    props.subject = "Structured faculty and staff directory extracted from the BRACU CSE website"
    props.author = "CampusConnect"
    props.keywords = "BRACU, CSE, faculty, staff, directory"

    doc.save(OUTPUT)


if __name__ == "__main__":
    items = extract_records()
    if not items:
        raise SystemExit("No directory records found")
    build_document(items)
    print(f"Created {OUTPUT}")
    print(f"Records: {len(items)}")
    counts = {}
    for item in items:
        counts[item["section"]] = counts.get(item["section"], 0) + 1
    for name, count in counts.items():
        print(f"{name}: {count}")
