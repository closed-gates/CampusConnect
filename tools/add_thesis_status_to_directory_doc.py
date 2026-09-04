from pathlib import Path
from collections import Counter

from lxml import html
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "thesis_supervisors_source.html"
DOC_PATH = ROOT / "artifacts" / "BRACU_CSE_Faculty_and_Staff_Directory.docx"
THESIS_URL = "https://cse.bracu.ac.bd/thesis/supervising/list"


def clean(value):
    return " ".join((value or "").split())


def extract_supervision_statuses():
    tree = html.fromstring(SOURCE.read_bytes())
    statuses = {}
    for card in tree.xpath("//div[contains(concat(' ', normalize-space(@class), ' '), ' fac-card ')]"):
        paragraphs = [clean(p.text_content()) for p in card.xpath(".//p")]
        email = next((p.lower() for p in paragraphs if "@" in p), "")
        status = next((p for p in paragraphs if p in {"Accepting", "Not Accepting"}), "")
        level_nodes = card.xpath("./div[contains(@class, 'absolute top-0 right-0')]/div/text()")
        level = clean(level_nodes[0]) if level_nodes else ""
        if email and status:
            statuses[email] = {"status": status, "level": level}
    return statuses


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=85, bottom=90, end=85):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def style_run(run, size=9.2, bold=False, color="222222"):
    run.font.name = "Aptos"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Aptos")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Aptos")
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def add_hyperlink(paragraph, text, url, color="0563C1"):
    rel_id = paragraph.part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    colour = OxmlElement("w:color")
    colour.set(qn("w:val"), color)
    props.append(colour)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_status_column(doc, statuses):
    matched = Counter()
    unmatched_emails = []
    widths = (0.38, 1.88, 1.36, 1.80, 1.78)

    for table in doc.tables:
        if len(table.columns) == 4:
            table.add_column(Inches(widths[-1]))
        if len(table.columns) != 5:
            continue

        for index, width in enumerate(widths):
            table.columns[index].width = Inches(width)

        header = table.rows[0]
        header.cells[4].text = "Thesis Supervision"
        for index, cell in enumerate(header.cells):
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_shading(cell, "176B73")
            set_cell_margins(cell, 105, 85, 105, 85)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if index in {0, 4} else WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                style_run(run, 9.0, True, "FFFFFF")

        for row_number, row in enumerate(table.rows[1:], 1):
            for index, cell in enumerate(row.cells):
                cell.width = Inches(widths[index])
            email = clean(row.cells[3].text).lower()
            info = statuses.get(email)
            status_cell = row.cells[4]
            status_cell.text = ""
            status_cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            status_cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_margins(status_cell)

            if info:
                label = info["status"]
                if info["level"]:
                    label += f"\n{info['level']}"
                color = "15803D" if info["status"] == "Accepting" else "B42318"
                fill = "EAF7EE" if info["status"] == "Accepting" else "FDECEC"
                matched[info["status"]] += 1
                style_run(status_cell.paragraphs[0].add_run(label), 8.8, True, color)
            else:
                fill = "F3F4F6"
                style_run(status_cell.paragraphs[0].add_run("Not listed"), 8.8, False, "6B7280")
                unmatched_emails.append(email)
            set_cell_shading(status_cell, fill)

    return matched, unmatched_emails


def update_intro_and_sources(doc, matched, total_statuses):
    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("This directory organizes"):
            paragraph.add_run(
                " The Thesis Supervision column reports whether a matching person is currently "
                "accepting thesis students and the published study level: U for undergraduate, "
                "P for postgraduate, and U & P for both. Not listed means the person does not "
                "appear on the thesis-supervisor page and is not treated as not accepting."
            )
            break

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Directory snapshot:"):
            paragraph.add_run(
                f"  |  Thesis matches: {sum(matched.values())} of {total_statuses} published supervisors"
            )
            break

    for paragraph in doc.paragraphs:
        if paragraph.text.startswith("Source:"):
            paragraph.add_run("  |  ")
            add_hyperlink(paragraph, "BRACU CSE Thesis Supervisors", THESIS_URL)
            break


def main():
    statuses = extract_supervision_statuses()
    doc = Document(DOC_PATH)
    matched, unmatched = add_status_column(doc, statuses)
    update_intro_and_sources(doc, matched, len(statuses))
    doc.core_properties.subject = (
        "Structured BRACU CSE faculty and staff directory with thesis supervision availability"
    )
    doc.save(DOC_PATH)

    directory_emails = {
        clean(row.cells[3].text).lower()
        for table in doc.tables
        for row in table.rows[1:]
        if len(row.cells) >= 4
    }
    source_only = sorted(set(statuses) - directory_emails)
    print(f"Updated {DOC_PATH}")
    print(f"Supervisor page entries: {len(statuses)}")
    print(f"Matched accepting: {matched['Accepting']}")
    print(f"Matched not accepting: {matched['Not Accepting']}")
    print(f"Directory entries not listed as supervisors: {len(unmatched)}")
    print(f"Supervisor entries absent from directory: {len(source_only)}")
    for email in source_only:
        print(f"SOURCE_ONLY {email}")


if __name__ == "__main__":
    main()
